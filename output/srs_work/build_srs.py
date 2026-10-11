from pathlib import Path
import math
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image, ImageDraw, ImageFont

# Small drawing adapter keeps diagram geometry independent of the installed runtime.
class Shape:
    def __init__(self,kind,xy,w,h,**kw):self.kind,self.xy,self.w,self.h,self.kw=kind,xy,w,h,kw
def Rectangle(xy,w,h,**kw):return Shape('rect',xy,w,h,**kw)
def Ellipse(xy,w,h,**kw):return Shape('ellipse',(xy[0]-w/2,xy[1]-h/2),w,h,**kw)
def Circle(xy,r,**kw):return Ellipse(xy,2*r,2*r,**kw)
class Ax:
    def __init__(self,w,h):self.w,self.h=w,h;self.im=Image.new('RGB',(int(w*200),int(h*200)),'white');self.draw=ImageDraw.Draw(self.im)
    def pt(self,x,y):return (int(x*200),int((self.h-y)*200))
    def set_xlim(self,*a):pass
    def set_ylim(self,*a):pass
    def axis(self,*a):pass
    def add_patch(self,s):
        x,y=s.xy;box=(*self.pt(x,y+s.h),*self.pt(x+s.w,y));fill=s.kw.get('fc',None);outline=s.kw.get('ec','black')
        if s.kind=='rect':self.draw.rectangle(box,fill=fill,outline=outline,width=2)
        else:self.draw.ellipse(box,fill=fill,outline=outline,width=2)
    def plot(self,xs,ys,**kw):self.draw.line([self.pt(x,y) for x,y in zip(xs,ys)],fill=kw.get('c','black'),width=max(1,int(kw.get('lw',1)*2)))
    def annotate(self,t,xy,xytext,**kw):self.plot([xytext[0],xy[0]],[xytext[1],xy[1]],c='#555555')
    def text(self,x,y,t,**kw):
        font=ImageFont.truetype('C:/Windows/Fonts/'+('calibrib.ttf' if kw.get('weight')=='bold' else 'calibri.ttf'),int(kw.get('fontsize',10)*1.35*200/72))
        pos=self.pt(x,y);box=self.draw.multiline_textbbox((0,0),t,font=font,spacing=7);w,h=box[2]-box[0],box[3]-box[1]
        px=pos[0]-(w/2 if kw.get('ha')=='center' else 0);py=pos[1]-(h/2 if kw.get('va')=='center' else 0)
        if kw.get('bbox'):self.draw.rectangle((px-4,py-4,px+w+4,py+h+10),fill='white')
        self.draw.multiline_text((px,py),t,font=font,fill='black',spacing=7,align='center' if kw.get('ha')=='center' else 'left')
    def savefig(self,path,**kw):self.im.save(path)
class Plot:
    def subplots(self,figsize):a=Ax(*figsize);return a,a
    def close(self,*a):pass
plt=Plot()

ROOT=Path(__file__).resolve().parent
OUT=ROOT.parent/'LABTRACK_SRS_v3.docx'
ROOT.mkdir(parents=True,exist_ok=True)

def usecase():
    fig,ax=plt.subplots(figsize=(11,7.5)); ax.set_xlim(0,11);ax.set_ylim(0,7.5);ax.axis('off')
    ax.add_patch(Rectangle((.3,.3),10.4,5.65,fill=False,lw=1.2));ax.text(.5,5.7,'LABTRACK system boundary',fontsize=11,weight='bold')
    lanes=[('Administrator',1.6,['UC02 Manage accounts','UC03 Configure labs','UC04 Maintain inventory','UC12 View analytics','UC11 Decide transfers']),('Lab assistant',4.2,['UC06 Decide reservations','UC07 Issue and return','UC08 Approve extensions','UC09 Decide event requests','UC11 Receive transfers']),('Faculty',6.8,['UC05 Reserve equipment','UC08 Request extension','UC09 Submit event request','UC10 Quick borrow','UC11 Request transfer']),('Student',9.4,['UC05 Reserve equipment','UC08 Request extension','View own loan history'])]
    for role,x,labels in lanes:
        ax.add_patch(Circle((x,7.02),.13,fill=False));ax.plot([x,x],[6.89,6.5],c='black');ax.plot([x-.25,x+.25],[6.73,6.73],c='black');ax.plot([x,x-.22],[6.5,6.22],c='black');ax.plot([x,x+.22],[6.5,6.22],c='black');ax.text(x,7.28,role,ha='center',fontsize=10)
        for j,label in enumerate(labels):
            y=5.1-j*.92;ax.plot([x,x],[6.2,y+.28],c='#bbbbbb',lw=.7,zorder=0)
            ax.add_patch(Ellipse((x,y),2.35,.63,fc='white',ec='black',lw=.8,zorder=2));ax.text(x,y,label.replace(' ', '\n',1) if len(label)>27 else label,ha='center',va='center',fontsize=8.5,zorder=3)
    fig.savefig(ROOT/'usecase.png',dpi=210,bbox_inches='tight');plt.close(fig)

def er(name,nodes,edges):
    fig,ax=plt.subplots(figsize=(11,8.1));ax.set_xlim(0,11);ax.set_ylim(0,8.1);ax.axis('off')
    for a,b,label in edges:
        xa,ya,_=nodes[a];xb,yb,_=nodes[b]
        dx,dy=xb-xa,yb-ya
        if abs(dx)>abs(dy): start=(xa+math.copysign(1.36,dx),ya);end=(xb-math.copysign(1.36,dx),yb)
        else:start=(xa,ya+math.copysign(.73,dy));end=(xb,yb-math.copysign(.73,dy))
        if xa==xb and abs(dy)>3:
            start=(xa-1.36,ya);end=(xb-1.36,yb)
            ax.plot([start[0],.15,.15,end[0]],[start[1],start[1],end[1],end[1]],c='#555555',lw=.9)
        else:
            ax.annotate('',xy=end,xytext=start,arrowprops=dict(arrowstyle='-',lw=.9,color='#555555'),zorder=0)
        mx,my=(start[0]+end[0])/2,(start[1]+end[1])/2
        if xa==xb and abs(dy)>3:mx,my=1.7,2.8
        ax.text(mx,my,label,ha='center',va='center',fontsize=8,bbox=dict(facecolor='white',edgecolor='none',pad=1.3))
    for title,(x,y,fields) in nodes.items():
        ax.add_patch(Rectangle((x-1.36,y-.73),2.72,1.46,fc='white',ec='black',lw=1))
        ax.add_patch(Rectangle((x-1.36,y+.36),2.72,.37,fc='#e9eef3',ec='black',lw=.7))
        ax.text(x,y+.54,title,ha='center',va='center',fontsize=9,weight='bold')
        ax.text(x-1.23,y+.22,'\n'.join(fields),ha='left',va='top',fontsize=8.5,linespacing=1.5)
    fig.savefig(ROOT/f'{name}.png',dpi=210,bbox_inches='tight');plt.close(fig)

usecase()
er('er_core',{
 'Department':(1.7,7.1,['PK id','UQ code','name']),
 'Lab':(5.5,7.1,['PK id','FK department_id','UQ code']),
 'EquipmentModel':(9.3,7.1,['PK id','FK lab_id','name, equipment_type']),
 'User':(1.7,4.2,['PK id','FK department_id optional','UQ email, university_id','role, account_status']),
 'Request':(5.5,4.2,['PK id','FK requester_id, model_id','FK lab_id, event_issue_id','quantity, kind, status']),
 'EquipmentUnit':(9.3,4.2,['PK asset_id','FK model_id, lab_id','status, condition']),
 'ExtensionRequest':(1.7,1.3,['PK id','FK request_id','FK requested_by_id','requested_due_date, status']),
 'Transaction':(5.5,1.3,['PK id','FK request_id, borrower_id','FK unit_asset_id, lab_id','due_date, status']),
 'TransferUnit':(9.3,1.3,['PK FK transfer_id','PK FK unit_asset_id'])
},[('Department','Lab','1 to 0..*'),('Lab','EquipmentModel','1 to 0..*'),('Department','User','0..1 to 0..*'),('User','Request','1 to 0..*'),('EquipmentModel','EquipmentUnit','1 to 0..*'),('Request','Transaction','1 to 0..*'),('Request','ExtensionRequest','1 to 0..*'),('EquipmentUnit','Transaction','1 to 0..*'),('EquipmentUnit','TransferUnit','1 to 0..*')])
er('er_support',{
 'Lab':(1.7,7.1,['PK id','FK department_id','name, code']),
 'LabAssistantAssignment':(5.5,7.1,['PK FK lab_id','FK assistant_id','FK assigned_by_id optional']),
 'User':(9.3,7.1,['PK id','role, account_status','name, email']),
 'InventoryImportBatch':(1.7,4.2,['PK id','FK lab_id','UQ lab_id plus content_hash']),
 'EventRequest':(5.5,4.2,['PK id','FK coordinator_id, lab_id','UQ FK event_issue_id optional','requested_items JSON, status']),
 'Notification':(9.3,4.2,['PK id','FK user_id','category, read, created_at']),
 'Transfer':(1.7,1.3,['PK id','FK from_lab_id, to_lab_id','FK requester_id','status, resolved_at']),
 'EventIssue':(5.5,1.3,['PK id','FK coordinator_id, lab_id','FK issued_by_id','due_date']),
 'NotificationEmail':(9.3,1.3,['PK FK notification_id','attempts, next_attempt_at','sent_at optional, failed'])
},[('Lab','LabAssistantAssignment','1 to 0..1'),('User','LabAssistantAssignment','1 to 0..*'),('Lab','InventoryImportBatch','1 to 0..*'),('User','Notification','1 to 0..*'),('EventRequest','EventIssue','0..1 to 0..1'),('Notification','NotificationEmail','1 to 0..1'),('Lab','Transfer','1 to 0..* per lab role')])

d=Document();sec=d.sections[0];sec.page_height=Inches(11.69);sec.page_width=Inches(8.27)
sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.72)
for n in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    s=d.styles[n];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.space_after=Pt(6)
d.styles['Normal'].paragraph_format.line_spacing=1.08
d.styles['Title'].font.size=Pt(30);d.styles['Heading 1'].font.size=Pt(18);d.styles['Heading 2'].font.size=Pt(12)
for n in ['Heading 1','Heading 2']:d.styles[n].paragraph_format.space_before=Pt(10);d.styles[n].paragraph_format.space_after=Pt(7)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('LABTRACK SRS 3.0  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for r in footer.runs:r.font.size=Pt(8)
d.core_properties.title='LABTRACK Software Requirements Specification';d.core_properties.author='LABTRACK Project Team';d.core_properties.subject='Laboratory inventory and equipment borrowing requirements'
def p(t='',style=None):return d.add_paragraph(t,style)
def h(t):d.add_heading(t,2)
def page(t):d.add_page_break();d.add_heading(t,1)
def table(headers,rows,widths=None):
    t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=widths or [6.83/len(headers)]*len(headers)
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for i,v in enumerate(headers):t.rows[0].cells[i].text=v
    rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
    for row in rows:
        for c,v in zip(t.add_row().cells,row):c.text=str(v)
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');pr.append(cant)
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');b.append(el)
            tc.append(b);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DFE7EF' if ri==0 else ('F5F7F9' if ri%2==0 else 'FFFFFF'));tc.append(sh)
            mar=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'90');el.set(qn('w:type'),'dxa');mar.append(el)
            tc.append(mar)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.0
                for r in para.runs:r.font.size=Pt(9);r.bold=ri==0
    p('')
def req(id,title,text,accept):
    h(f'{id} {title}');p(text);q=p();q.add_run('Acceptance ').bold=True;q.add_run(accept)
def image(name,caption):
    d.add_picture(str(ROOT/name),width=Inches(6.8));p(caption,'Caption')

p('LABTRACK', 'Title');p('Software Requirements Specification','Subtitle');p('Laboratory Inventory and Equipment Borrowing System')
p('Version 3.0');p('08 October 2026')
p('This specification defines the required behavior of LABTRACK for administrators, laboratory assistants, faculty members, and students. It establishes the scope, interfaces, business rules, data relationships, quality requirements, and acceptance criteria used to review and validate the system.')
table(['Document control','Value'],[['Document ID','LT SRS 003'],['Prepared for','LABTRACK project stakeholders and evaluators'],['Prepared by','LABTRACK Project Team'],['Status','For stakeholder review and approval'],['Structure','IEEE style SRS informed by ISO/IEC/IEEE 29148:2018'],['Baseline','Current project workflows and database schema as of 08 October 2026']],[1.65,5.18])
h('Approval record');table(['Responsibility','Name','Approval and date'],[['Project owner','To be assigned','Pending'],['Technical reviewer','To be assigned','Pending'],['Academic supervisor or evaluator','To be assigned','Pending']],[2.6,1.6,2.63])
p('Approval establishes the requirements baseline. Proposed quality targets and unresolved policy decisions must be confirmed before acceptance testing is finalized.')

page('Document navigation')
toc=[('1 Introduction',3),('2 Overall description',4),('3 External interface requirements',5),('4 User roles and use case model',6),('5 Functional requirements',7),('5 Functional requirements continued',8),('5 Functional requirements continued',9),('5 Functional requirements continued',10),('6 Use case specifications',11),('6 Use case specifications continued',12),('7 Data requirements and core ER model',13),('7 Supporting ER model',14),('7 Data dictionary and integrity rules',15),('8 Nonfunctional requirements',16),('9 Verification and traceability',17),('10 Assumptions and open decisions',18),('Appendix A References and glossary',19)]
table(['Section','Page'],toc,[5.9,.93])
h('Revision history');table(['Version','Date','Change'],[['2.0','18 August 2026','Existing project specification'],['3.0','08 October 2026','Restructured SRS with traceable requirements, use cases, ER models, acceptance criteria, and policy decisions']],[.7,1.4,4.73])
p('Reading guide: stakeholders should review scope and roles first, developers should review functional and data requirements, and testers should use the acceptance criteria and traceability matrix. Requirement identifiers remain stable across future revisions.')

page('1 Introduction')
h('1 1 Purpose');p('LABTRACK shall provide a centralized web application for managing laboratory equipment and accountable borrowing. This SRS specifies externally observable behavior and essential data constraints so stakeholders can agree on what the system must do and how it will be accepted.')
h('1 2 Scope');p('The system covers institutional registration, role based access, departments and labs, assistant assignments, equipment models and physical units, QR identification, inventory import, reservations, approvals, issue and return operations, extensions, faculty event loans, quick borrowing, inter-lab transfers, notifications, analytics, and administrative borrowing settings.')
p('LABTRACK records individual physical assets and their circulation history. Standard reservation approval permits collection at the counter; faculty event approval allocates and issues the selected units as one operation. These workflows are intentionally specified separately.')
h('1 3 Exclusions');p('The baseline does not include payments, fines collection, supplier ordering, purchase order persistence, external student information system synchronization, biometric authentication, RFID hardware, or a native mobile application. Procurement suggestions are advisory calculations. A procurement screen action is not evidence of a completed purchasing workflow.')
h('1 4 Intended audience');p('The audience includes the project owner, academic evaluators, developers, testers, system administrators, and representatives of each operational role. This document supports requirements review, implementation planning, and acceptance testing.')
h('1 5 Specification conventions');p('“Shall” denotes a requirement. FR identifies a functional requirement; NFR identifies a quality or operational requirement; BR identifies a business rule; UC identifies a use case; and AT identifies an acceptance test. All FR requirements are mandatory within the baseline scope. NFR targets explicitly marked proposed require stakeholder approval and are not measured performance claims.')
h('1 6 Standards basis');p('The organization follows the familiar SRS sequence of introduction, overall description, interfaces, specific requirements, and supporting models. ISO/IEC/IEEE 29148:2018 is the requirements engineering reference [R1]. This document does not assert certification or full clause by clause conformity.')

page('2 Overall description')
h('2 1 Product perspective');p('LABTRACK is a browser based client server system. A React interface communicates with a FastAPI service through an HTTP API. PostgreSQL stores authoritative records. SQLAlchemy provides persistence and Alembic controls schema migrations. Gmail API or SMTP may deliver optional email notifications.')
h('2 2 Product functions');table(['Capability','Operational result'],[['Identity and access','Institutional accounts and role scoped operations'],['Inventory','Model catalog, unique assets, QR labels, and stock status'],['Circulation','Reservation decisions, counter issue, return, and extension'],['Events and daily loans','Faculty event equipment and same day quick loans'],['Transfers','Controlled movement of assets between labs'],['Oversight','History, notifications, reports, suggestions, and policies']],[1.65,5.18])
h('2 3 Operating environment');p('The frontend operates in a modern desktop or mobile web browser with network access to the API. Camera scanning requires camera permission and a supported secure browser context; manual asset entry provides a fallback. Deployment requires a reachable PostgreSQL instance and the configured backend service. A reverse proxy and Docker Compose are optional deployment arrangements.')
h('2 4 Design and implementation constraints');p('The present implementation uses React 19 with Vite, Python FastAPI, JWT authentication, PostgreSQL, SQLAlchemy, and Alembic. Browser storage shall not replace the database as the authoritative inventory or circulation store. Secrets and provider credentials are deployment configuration and shall not appear in public source or client bundles.')
h('2 5 Dependencies');p('Institutional departments and lab assignments must be maintained by an administrator. Physical handover and condition inspection depend on the lab assistant. Email depends on enabled provider credentials and connectivity. A QR code identifies an asset; it does not independently prove user authorization or physical custody.')
h('2 6 User characteristics');p('Students and faculty require basic browser skills. Assistants require training in scan verification, handover, condition inspection, and transfer receipt. Administrators require knowledge of institutional roles, lab assignments, and borrowing policies.')

page('3 External interface requirements')
req('IF01','User interface','The system shall provide login and registration screens, role specific dashboards, inventory views, requests, loan history, notifications, and the management screens permitted for each role.','Each role reaches its dashboard and authorized screens; direct navigation to a forbidden screen cannot grant an API operation.')
req('IF02','API interface','The frontend shall use the configured API base path, normally /api/v1. Structured application data shall use JSON; login shall support the backend credential form. Authenticated requests shall carry the issued bearer token.','Valid requests succeed; missing or invalid authentication and unauthorized roles receive appropriate non-success responses.')
req('IF03','Scanning and printing','The system shall support equipment QR lookup and readable asset identifiers. Counter workflows shall permit scanner or manual input. Generated labels shall associate each code with its physical unit.','A printed label resolves to the correct asset, and an unreadable label can be handled by entering the asset ID.')
req('IF04','Inventory import','The system shall accept the equipment import data contract with a selected target lab, equipment details, quantities, and optional unit information. Spreadsheet parsing and the JSON import API shall be treated as separate interface steps.','Malformed data is rejected with a useful error; repeating a successful identical batch for the same lab adds no stock.')
req('IF05','Notification delivery','Portal notifications shall be available without an email provider. Optional Gmail API or SMTP delivery shall use configured credentials and the persistent email outbox.','Portal operation remains usable with email disabled; enabled delivery records its attempts and result.')
req('IF06','Dates and infrastructure','The user interface shall display calendar dates as dd-mm-yyyy. Stored transaction timestamps shall be timezone aware. Quick loans shall use the institutional Asia/Kolkata end of day. The database shall be prepared through migrations.','A local same day due time corresponds to 23:59:59 in the institutional timezone; API health can be checked at /health.')

page('4 User roles and use case model')
image('usecase.png','Figure 1  UML use case overview of the principal role associations')
p('Repeated UC identifiers represent the same use case shared by multiple roles. Authentication, catalog browsing, personal notifications, and profile access are common supporting use cases and are omitted from this overview for readability. Association lines indicate participation, not the order of execution.')
table(['Role','Scope and authority'],[['Administrator','Manage accounts, departments, labs, assignments, models, inventory, settings, reports, and transfer decisions.'],['Lab assistant','Operate assigned labs; decide reservations and events; issue, return, approve extensions, and receive transfers.'],['Faculty','Browse, reserve, view own history, request extensions and transfers, coordinate event requests, and receive quick loans.'],['Student','Browse, reserve, view own history, and request extensions.']],[1.25,5.58])

page('5 Functional requirements')
p('Identity and organizational management. Related use cases: UC01 to UC03.')
req('FR01','Registration','The system shall register students, faculty, and assistants using name, institutional email, credentials, institutional identifier where required, and a valid selected department. It shall reject public administrator registration.','Duplicate email or institutional ID is rejected; invalid department or administrator role cannot create an account.')
req('FR02','Account eligibility','The system shall normalize emails to lowercase and institutional IDs to uppercase. Students shall use charusat.edu.in and faculty or assistants shall use charusat.ac.in, subject only to recorded legacy exemptions. New student accounts shall be ACTIVE; faculty and assistant registrations shall be PENDING.','A pending or deactivated account cannot authenticate as an active user; department display follows the chosen database department.')
req('FR03','Authentication and authorization','The system shall verify credentials, issue authenticated sessions, and enforce role and account status on the backend. Assistants shall be restricted to their assigned labs for operational actions.','A forbidden API request fails even if the caller bypasses the frontend; valid credentials for an active account permit its authorized operations.')
req('FR04','Account administration','Administrators shall review staff account status, maintain user details, and deactivate accounts while retaining referenced borrowing history. User directories shall distinguish staff and students.','An authorized change is reflected in the directory; deactivation prevents access and leaves historical borrower references intact.')
req('FR05','Departments labs and assignments','Administrators shall maintain departments and labs with unique codes and assign an assistant to a lab. One lab shall have at most one assignment; one assistant may manage multiple labs.','A second assignment cannot coexist for one lab; operational lab visibility updates when an assignment changes.')

page('5 Functional requirements continued')
p('Inventory and reservations. Related use cases: UC04 to UC06.')
req('FR06','Equipment models and units','Administrators shall maintain model details including name, category, owning catalog lab, description, and equipment type. Administrators and authorized assistants shall add units with unique asset IDs and current lab ownership.','One model supports multiple physical units; a unit retains its own asset ID, status, optional serial number, and condition.')
req('FR07','Asset identification','The system shall allocate collision free asset IDs using LT followed by lab code, category code, and a sequence, and support QR labels for those IDs.','Concurrent allocations produce distinct IDs; a scan retrieves the corresponding unit, not merely a model.')
req('FR08','Inventory import','Administrators and assigned assistants shall import validated equipment batches into an authorized target lab. A successful batch shall be identified by lab and content hash to prevent duplicate additions.','Submitting the same successful content twice to the same lab does not increase stock the second time; invalid rows do not produce a partial successful batch.')
req('FR09','Catalog and availability','Authenticated users shall browse equipment models and view unit availability. Inventory operations shall distinguish AVAILABLE, ISSUED, MAINTENANCE, and IN_TRANSIT units.','Issued, maintenance, or in transit units cannot be selected as available for handover; current ownership is respected after a transfer.')
req('FR10','Reservation submission','Students and faculty shall submit a standard request for a model and lab with positive quantity, purpose, start time, and later end time, within applicable borrowing limits.','Valid submission creates a PENDING request; invalid quantity, date order, or policy limit is rejected.')
req('FR11','Reservation decision','Only an assistant assigned to the request lab shall approve or reject a PENDING standard request. Rejection shall record a reason and the decision shall retain actor and time.','Approval changes the request to APPROVED without issuing a unit; repeated decisions on an already decided request fail.')

page('5 Functional requirements continued')
p('Counter circulation and extensions. Related use cases: UC07, UC08, and UC10.')
req('FR12','Approved checkout','An assigned assistant shall identify the borrower and approved request, validate an available scanned or entered unit against model and current lab, and confirm handover. Each issued unit shall create one ACTIVE transaction.','Wrong model, wrong lab, or unavailable unit is rejected. A multi-unit request becomes ISSUED when its requested quantity has been issued.')
req('FR13','Walk in issue','An assigned assistant shall directly issue an available standard unit to an eligible student or faculty borrower with a valid future due date and applicable duration limits.','The operation creates an ISSUED WALK_IN request and its ACTIVE unit transaction without a prior reservation.')
req('FR14','Return and condition','An assigned assistant shall locate the active unit loan, record return time and condition remarks, close the transaction, and update the unit condition. Damage or repair indications shall route the unit to MAINTENANCE; otherwise it shall become AVAILABLE.','A repeated return cannot close a loan twice. A request becomes RETURNED only when none of its unit transactions remains active.')
req('FR15','Borrowing history and overdue state','The system shall show role scoped requests and transactions, including borrower, asset, issue time, due time, return time, and condition. Overdue shall mean an ACTIVE transaction with due_date earlier than the current time.','Returned transactions are excluded from overdue counts; students and faculty see their own records and assistants see assigned lab records.')
req('FR16','Extension request and approval','A student or faculty borrower shall request a later due date for an own ISSUED request with an active loan. At most one PENDING extension shall exist per request. Assigned assistants shall approve a pending extension and update all active transactions for that request.','Approval records the decision, changes the due dates, and increments reissue counts without another physical scan. Nonowners and duplicate pending extensions are rejected.')
req('FR17','Faculty quick borrowing','Assigned assistants shall issue QUICK_BORROW units to faculty through a direct counter flow. The due time shall be the end of that institutional day. Quick return shall close the related loan.','Student quick borrowing, a standard equipment type, or a nonavailable unit is rejected; faculty loan due time is 23:59:59 Asia/Kolkata.')

page('5 Functional requirements continued')
p('Events, transfers, and oversight. Related use cases: UC09, UC11, and UC12.')
req('FR18','Faculty event requests','Faculty shall submit an event or club request with event name, purpose, target lab, future due date, and requested models and quantities or an explicit supported asset list. Only the assigned assistant shall approve or reject it. Approval shall validate the allocation and atomically create an event issue and related loans.','A failed allocation issues no subset. Successful approval links the event request to its event issue and issues all selected available units; rejection records a reason.')
req('FR19','Transfer lifecycle','Faculty, assistants, and administrators shall request movement of available units between distinct labs. Only administrators shall approve or reject a PENDING transfer. Approval shall mark units IN_TRANSIT. An assistant assigned to the destination shall complete receipt and change current ownership.','Completion makes the units AVAILABLE in the destination lab. Only a pending transfer may be cancelled by its requester or an administrator; conflicting active transfers are rejected.')
req('FR20','Notifications','The system shall create recipient scoped portal notifications for supported circulation and system events and allow the recipient to mark them read. Enabled email delivery shall use the outbox.','One user cannot read or alter another user’s personal notification; mail failure does not change the equipment custody state.')
req('FR21','Reports and procurement suggestions','Administrators shall view system totals, borrowing trends, lab utilization, frequently borrowed models, and advisory procurement suggestions based on inventory and pending demand.','Results reconcile with persisted records. Suggestions show their reason and stock context; they do not create an actual supplier order or assert machine learning predictions.')
req('FR22','Borrowing settings','Administrators shall maintain supported system settings, including borrowing durations and student item limits. Effective limits shall be visible to operational users.','A nonadministrator cannot alter settings. Reservation validation uses stored duration and item limits; extension policy consistency is tracked as decision OD01.')

page('6 Use case specifications')
p('UC01 Authenticate: an active user submits valid credentials and receives access to the role dashboard. Invalid credentials, pending approval, rejection, or deactivation prevent entry. UC02 Manage accounts, UC03 Configure labs, and UC04 Maintain inventory are governed by FR04 to FR08.')
h('UC05 Submit equipment reservation');p('Primary actor: Student or Faculty. Preconditions: active account, selectable standard model, and target lab. Trigger: the borrower selects Request Equipment. Main flow: select the model, enter quantity and purpose, choose required dates, review limits, and submit. The system validates the request and stores PENDING status. Alternative flow: invalid quantity, dates, or a breached limit returns an actionable validation message. Postcondition: the request is visible to the borrower and the assigned assistant. Requirements: FR09 and FR10.')
h('UC06 Decide reservation');p('Primary actor: assigned Lab assistant. Preconditions: a PENDING request in an assigned lab. Main flow: open the queue, inspect details, and approve or reject with a reason. The system records decision and status and notifies the borrower. Alternatives: an out of scope lab or already decided request is rejected. Postcondition: APPROVED or REJECTED request; no physical custody change on approval. Requirements: FR11 and FR20.')
h('UC07 Issue and return equipment');p('Primary actor: assigned Lab assistant; supporting actor: borrower. Preconditions for standard issue: approved request and available unit matching its model and lab. Main flow: identify the borrower, select the request, scan or enter the asset ID, verify the displayed unit, and confirm handover. The system creates an ACTIVE transaction and updates unit and request status. Repeat for additional units until quantity is met. Alternatives: mismatched or unavailable assets are rejected; a walk in uses the direct issue flow and due date validation.')
p('Return flow: identify the asset, retrieve its active loan, inspect condition, enter remarks, and confirm receipt. The system records return time, closes the transaction, and routes the unit to AVAILABLE or MAINTENANCE. The request is RETURNED only after all active loans close. Alternatives: missing active loan or unassigned lab prevents the operation. Postcondition: recorded custody history agrees with unit status. Requirements: FR12 to FR15.')
h('UC08 Extend a loan');p('Primary actors: Student or Faculty, then assigned Lab assistant. Preconditions: own ISSUED request, active transactions, and no pending extension. Main flow: borrower selects a later permitted due date and supplies a reason; the system stores a PENDING extension; assistant approves; the system updates all active transaction due dates and reissue counts. Alternatives: ownership, date, limit, and duplicate checks fail without changing existing due dates. Postcondition: approval is recorded without physical reissue. Requirement: FR16.')

page('6 Use case specifications continued')
h('UC09 Request and issue event equipment');p('Primary actors: Faculty coordinator and assigned Lab assistant. Preconditions: active faculty account and a selected lab with required equipment. Main flow: coordinator enters event name, purpose, due date, and model quantities; submits request; assistant reviews and selects specific available units matching each model quantity; assistant approves. The system creates the EventIssue, EVENT requests, and unit transactions, and marks the event request APPROVED. Alternatives: reject with reason, mismatched allocation, unavailable unit, or invalid date. A validation failure leaves the request pending and creates no partial issue. Postcondition: all issued units are attributable to the coordinator and event. Requirement: FR18.')
h('UC10 Borrow a faculty quick loan');p('Primary actor: assigned Lab assistant; beneficiary: Faculty. Preconditions: eligible faculty borrower and available QUICK_BORROW unit. Main flow: identify faculty, scan unit, verify lab assignment and type, confirm issue, and later record quick return. Alternatives: student borrower or standard equipment is rejected. Postcondition: same day loan is recorded and then closed on return. Requirement: FR17.')
h('UC11 Transfer equipment between labs');p('Primary actors: requester, Administrator, and destination Lab assistant. Preconditions: distinct source and destination labs, available source units, and no conflicting active transfer. Main flow: requester selects units and reason; administrator approves; system marks the units IN_TRANSIT; destination assistant confirms receipt; system changes unit lab_id and returns units to AVAILABLE. Alternatives: administrator rejects while pending; requester or administrator cancels while pending; unauthorized receipt or invalid state change is rejected. Postcondition: asset identities and historical loans remain intact after relocation. Requirement: FR19.')
h('UC12 Review operations');p('Primary actor: Administrator. Preconditions: active administrator account. Main flow: open reports, inspect stock and circulation counts, review monthly trends and model use, and inspect procurement suggestions. Alternative: empty datasets show zero or no suggestions without invented records. Postcondition: administrator can make an informed inventory decision; no actual order is placed. Requirement: FR21.')
h('Lifecycle summary');table(['Record','Permitted progression'],[['Standard request','PENDING to APPROVED or REJECTED; APPROVED to ISSUED as quantity is fulfilled; ISSUED to RETURNED when all loans close.'],['Event request','PENDING to APPROVED with linked issue, or REJECTED with reason.'],['Transaction','ACTIVE to RETURNED; overdue is derived, not a stored status.'],['Transfer','PENDING to APPROVED, REJECTED, or CANCELLED; APPROVED to COMPLETED.']],[1.4,5.43])
p('CANCELLED exists in the standard request schema, but a borrower cancellation workflow is not part of this baseline. The partial issue and early return case requires the decision recorded as OD04.')

page('7 Data requirements and core ER model')
image('er_core.png','Figure 2  Core inventory and circulation ER model with selected keys')
p('PK means primary key; FK means foreign key; UQ means unique. A label reads from the first entity to the second: 1 to 0..* means one parent can have zero or many child records. Each child references one parent unless the relationship is marked optional.')
p('This view shows the principal relationships, not every foreign key. Lab is also referenced by EquipmentUnit, Request, and Transaction. User is also referenced as borrower, requester, decision maker, issue operator, and receiving operator. Request.model_id references EquipmentModel. Complete key fields are summarized in the data dictionary.')
h('7 1 Catalog and ownership');p('EquipmentModel.lab_id is the catalog lab; EquipmentUnit.lab_id is current physical ownership and can change during transfer. Historical Request.lab_id and Transaction.lab_id preserve the operating lab of those records. Asset IDs remain stable when ownership changes.')
h('7 2 Loan cardinality');p('A request may have zero or many unit transactions. A unit may have many historical transactions but at most one ACTIVE transaction. A request may have many historical extensions but at most one PENDING extension. TransferUnit connects individual assets to transfers using a composite primary key.')

page('7 Supporting ER model')
image('er_support.png','Figure 3  Supporting organization event transfer and notification ER model')
p('The optional EventRequest.event_issue_id is unique: a request links to at most one issue and an issue links back to at most one request. Approval requires that link. NotificationEmail is an optional outbox record keyed by its Notification. A lab has at most one assistant assignment, while one assistant may have multiple lab assignments.')
p('Additional relationships omitted from the image: EventRequest and EventIssue each reference Lab and their coordinator User. EventIssue is referenced by EVENT Request rows. Transfer references Lab twice, as source and destination, and relates to TransferUnit in Figure 2. User references also identify administrative and operational decision makers.')
h('7 3 Nonrelational event selections');p('EventRequest.requested_items and unit_asset_ids are JSON values. They are validated by the application rather than represented as child tables with asset foreign keys. They are not shown as FK relationships in either figure.')
h('7 4 Independent supporting entities');p('AssetSequence stores one atomic counter per asset prefix. SystemSetting stores unique configuration keys. LegacyColumnValue preserves values migrated from historical columns. These entities do not introduce application foreign key relationships and are listed in the dictionary.')

page('7 Data dictionary and integrity rules')
table(['Entity','Identity and important fields'],[
 ['Department','id PK; code UQ; name, hod_name'],['User','id PK; email and university_id UQ; department_id optional FK; role, account_status, hashed_password; approval and activation metadata'],['Lab','id PK; department_id FK; code UQ; name, location'],['LabAssistantAssignment','lab_id PK FK; assistant_id FK; assigned_by_id optional FK; assigned_at'],['EquipmentModel','id PK; lab_id FK; name, category, equipment_type, description'],['EquipmentUnit','asset_id PK; model_id and lab_id FK; serial_number, status, condition, qr_code_url'],['Request','id PK; requester_id, model_id, lab_id FK; event_issue_id optional FK; kind, quantity, dates, status, decision metadata'],['Transaction','id PK; request_id, unit_asset_id, borrower_id, lab_id FK; issue, due and return dates; status, condition and operator metadata'],['ExtensionRequest','id PK; request_id and requested_by_id FK; requested_due_date, status, decision metadata'],['EventRequest','id PK; coordinator_id and lab_id FK; event_issue_id optional unique FK; JSON selections, status, due_date'],['EventIssue','id PK; coordinator_id, lab_id, issued_by_id FK; event_name, purpose, issue_date, due_date'],['Transfer and TransferUnit','Transfer id PK; source, destination and user FKs; status. TransferUnit composite PK of transfer_id and unit_asset_id, both FKs'],['Notification and NotificationEmail','Notification id PK; user_id FK; message, category, read. Outbox notification_id PK FK; attempts, scheduling and outcome'],['AssetSequence','prefix PK; last_value nonnegative'],['InventoryImportBatch','id PK; lab_id FK; lab plus content_hash unique; result JSON'],['SystemSetting','id PK; key UQ; value'],['LegacyColumnValue','id PK; table_name, row_key, column_name, value, migrated_at']],[1.68,5.15])
p('BR01 Positive quantities and valid date order shall be enforced. BR02 A unit shall have at most one active loan. BR03 A request shall have at most one pending extension. BR04 Rejected requests shall retain a reason. BR05 Returned transactions shall have a return timestamp; active transactions shall not. BR06 Transfers shall use distinct labs. BR07 Referenced circulation history shall be retained. BR08 Duplicate successful inventory batches shall not add stock again.')

page('8 Nonfunctional requirements')
p('Security and integrity requirements are mandatory. Numeric performance and operational service targets below are proposed acceptance targets for stakeholder approval; they are not claims about measurements already achieved.')
req('NFR01','Access and credentials','The system shall enforce account status, role, resource ownership, and lab assignment on the server; store passwords as hashes; and keep deployment secrets outside public source and client output.','Inspect credential storage and run allowed and denied API scenarios for every role. Production transport shall use HTTPS.')
req('NFR02','Transactional integrity','Custody and inventory mutations shall preserve BR01 to BR08 under failure and competing operations. Multi-unit event issuance shall succeed or fail as a complete operation.','Concurrent attempts to issue the same unit create only one active loan; a failed event allocation changes no subset of its units.')
req('NFR03','Response time and capacity','Proposed target: at 50 concurrent users with 10,000 units and 100,000 historical transactions, 95 percent of catalog, request-list, and asset-lookup API responses shall complete within 2 seconds; checkout and return within 3 seconds, excluding user scanning time and external email.','Measure under a documented representative server configuration and report dataset, concurrency, p95 timing, and errors.')
req('NFR04','Availability and recovery','Proposed target: 99 percent availability during scheduled institutional operating hours, excluding announced maintenance. Operations shall support daily database backups with a recovery point of 24 hours and recovery time of 4 hours.','Review uptime records and restore a backup into an isolated environment; reconcile asset and active loan counts.')
req('NFR05','Usability and accessibility','The interface shall provide labeled controls, clear validation and success messages, readable statuses, and manual scan fallback. Proposed acceptance: core tasks work by keyboard and on supported current desktop and mobile browsers.','A representative user completes reservation, checkout, and return; verify keyboard focus, readable forms, and denied camera fallback.')
req('NFR06','Maintainability and portability','The system shall separate frontend, API, and persistence concerns, provide documented setup and configuration, and apply schema changes through versioned migrations.','A fresh configured database reaches the current schema; frontend production build succeeds; documented health and API endpoints respond.')
req('NFR07','Privacy and auditability','Personal data and operational history shall be exposed only within authorized scope. Important decisions shall retain the actor and timestamp where supported by the model. Proposed retention policy requires OD02 approval.','Verify role filtered record access and inspect decision metadata; review missing custody operator coverage under OD03.')

page('9 Verification and traceability')
p('Acceptance requires reviewed requirements, passing functional scenarios, resolved blocking policy decisions, and evidence for approved quality targets. The following matrix defines planned verification; it does not claim that these tests have already passed.')
table(['Test','Requirements and use cases','Expected evidence'],[
 ['AT01','FR01 to FR04; UC01 UC02; NFR01','Register each role; reject duplicates and public admin; deny inactive and invalid users.'],['AT02','FR05; UC03; NFR01','Assign labs; demonstrate one assistant per lab and deny operations outside assignment.'],['AT03','FR06 to FR09; UC04; BR08','Create models and unique units; scan labels; repeat import with no duplicate stock.'],['AT04','FR10 FR11; UC05 UC06','Validate limits and dates; only assigned assistant decides; approval does not issue.'],['AT05','FR12 FR13; UC07; BR02','Issue valid asset; reject wrong model or lab; create quantity bounded transactions.'],['AT06','FR14 FR15; UC07; BR05','Return each unit once; route damaged units to maintenance; reconcile overdue results.'],['AT07','FR16; UC08; BR03','Deny nonowner and duplicate pending extension; update all active due dates on approval.'],['AT08','FR17; UC10','Faculty only, quick equipment only; verify local end of day and quick return.'],['AT09','FR18; UC09; NFR02','Approve correct event allocation; force one invalid unit and confirm zero partial issue.'],['AT10','FR19; UC11; BR06','Test request, admin approval, destination receipt, conflict, and pending cancellation.'],['AT11','FR20; IF05','Test recipient scope, read state, disabled email, delivery failure, and outbox attempts.'],['AT12','FR21 FR22; UC12','Reconcile reports and suggestions; deny nonadmin reports or setting mutations.'],['AT13','NFR02; BR02 BR03','Competing unit issuance and extension submissions preserve uniqueness and state.'],['AT14','NFR03 NFR04','Load report for approved targets and backup restore evidence.'],['AT15','IF01 to IF04 IF06; NFR05 NFR06','Browser and keyboard tasks, manual scan fallback, date display, migration and build checks.'],['AT16','NFR07; BR07','User scope and historical preservation; inspect custody and decision metadata.']],[.65,2.35,3.83])
h('Acceptance evidence');p('Each test record shall identify its environment, data, actor role, steps, expected result, actual result, and evidence. Failures affecting authorization, duplicate custody, partial event issue, or historical integrity shall block acceptance. Quality targets shall be tested only after their workloads and thresholds are approved.')

page('10 Assumptions and open decisions')
h('10 1 Assumptions');p('The institution supplies valid department data, responsible lab assistants, and borrower identities. Equipment is physically labeled and assistants confirm actual handover. Core transactions require connectivity; offline custody changes are outside scope. Optional email provider failure does not remove portal access. Administrators operate backup and retention processes.')
h('10 2 Default borrowing policy');p('Current startup defaults are 14 days for students, 30 days for faculty, and a student item limit of 3. Quick borrowing is faculty only and due on the same institutional day. The selected registration department determines the portal label. Supported registration departments are CE, CSE, IT, AIML, and EC. These are institutional defaults, not universal laboratory rules.')
h('10 3 Decisions requiring confirmation');table(['ID','Decision and consequence','Owner'],[
 ['OD01','Reservation durations are configurable, but extension validation currently uses fixed 14 and 30 day limits. Confirm one authoritative policy, precise elapsed-time boundary, and renewal eligibility before changing limits.','Project owner'],['OD02','Approve retention period for user, transaction, notification, and email records; backup storage and restore responsibility; no legal retention period is assumed.','Institution administrator'],['OD03','Confirm whether every checkout and quick return must record its operator. Fields exist, but actor recording is not uniform across present flows.','Technical lead'],['OD04','Define partial checkout followed by an early return: remaining entitlement and request status must not permit overissue or lose uncollected quantities.','Lab operations owner'],['OD05','Approve proposed load, response time, operating availability, backup, and recovery targets and test hardware.','Project owner'],['OD06','Confirm whether current unit lab or original model lab governs lab utilization reporting after transfer. Custody uses current unit ownership.','Lab operations owner'],['OD07','Confirm cancellation and rejection capabilities for extensions and standard reservations before adding them; schema states alone do not establish a complete user workflow.','Project owner']],[.6,4.9,1.33])
h('10 4 Change control');p('A proposed change shall identify affected requirement IDs, actors, interfaces, data constraints, and tests. The project owner and technical reviewer shall approve its scope and update the version and revision history. Approval of this SRS does not certify that the implementation already meets every requirement.')

page('Appendix A References and glossary')
h('References');p('[R1] ISO/IEC/IEEE 29148:2018, Systems and software engineering — Life cycle processes — Requirements engineering. Official overview: https://www.iso.org/standard/72089.html. Used as the requirements engineering reference; full formal conformity is not asserted.')
p('[R2] LABTRACK README.md. Project scope, role descriptions, environment, setup, and feature overview.')
p('[R3] backend/app/db/models.py and backend/alembic/versions. Canonical entities, keys, statuses, database constraints, and event request item migration.')
p('[R4] backend/app/api/endpoints/auth.py, inventory.py, borrowing.py, events.py, transfers.py, labs.py, notifications.py, analytics.py, and settings.py. Operational authorization and workflow behavior.')
p('[R5] src/App.jsx and src/pages. User routes and role specific interfaces. src/pages/analytics/SmartProcurementPage.jsx establishes that the current purchase order action is a notification rather than a persisted order.')
p('[R6] backend/app/startup.py and backend/app/services. Policy defaults, institutional identity rules, lab access, asset counters, and email delivery services.')
p('[R7] LabTrack_SRS_v2.docx. Previous specification retained as revision context; current repository behavior governs this baseline.')
h('Glossary');table(['Term','Meaning'],[['Asset','One individually identifiable physical equipment unit.'],['Model','A catalog entry describing an equipment type with multiple units.'],['Reservation','A request for equipment before physical checkout.'],['Transaction','The circulation record for one issued unit.'],['Overdue','An active loan whose due timestamp is earlier than now.'],['Event issue','A grouped equipment issue associated with a faculty event coordinator.'],['Quick borrow','Faculty counter loan for eligible equipment, due the same day.'],['In transit','Unit reserved by an approved inter-lab transfer awaiting receipt.'],['JWT','JSON Web Token used for authenticated API requests.'],['ER model','Entity relationship representation of data and cardinalities.'],['Outbox','Persistent record used to schedule and track notification email delivery.']],[1.4,5.43])

for element in d.styles.element.xpath('.//w:pBdr'):
    element.getparent().remove(element)
for element in d.element.xpath('.//w:pBdr'):
    element.getparent().remove(element)
in_nfr=False
for para in d.paragraphs:
    if para.text=='8 Nonfunctional requirements':in_nfr=True
    if para.text=='9 Verification and traceability':in_nfr=False
    if in_nfr and para.style.name=='Normal':
        para.paragraph_format.space_after=Pt(4)
        para.paragraph_format.line_spacing=1.0
        for run in para.runs:run.font.size=Pt(10)
d.save(OUT)
print(OUT)
